"""Synthetic examples only; no enterprise records or external files."""
import io
from openpyxl import Workbook

HEADERS = ['STT','Tên VTTB','Mã VTTB','ĐVT','Chủng loại','Loại tài sản','Serial','Số lượng','Đơn vị quản lý','Phân loại thiết bị','', 'Mã đơn vị','Chi nhánh cũ','Chi nhánh mới','Bưu cục','Đơn vị quản lý','Mã nhân viên','Nhân viên quản lý','Ngày đưa vào sử dụng','Tình trạng']


def synthetic_files(rows=1000):
    if not 1 <= rows <= 250000:
        raise ValueError('Số dòng phải từ 1 đến 250.000.')
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('Sheet1')
    sheet.append(HEADERS)
    for i in range(rows):
        code = i % 200 + 1
        name = f'Thiết bị mẫu {code:04d}' if i % 83 else 'Tên mẫu chưa phân loại'
        kind = 'Vật tư' if i % 7 == 0 else 'Thiết bị'
        status = 'Hỏng' if i % 29 == 0 else 'Tốt' if i % 11 == 0 else 'Đang sử dụng'
        used = '' if i % 13 == 0 else '31/02/2025' if i % 97 == 0 else '01/01/2099' if i % 101 == 0 else f'{i % 28 + 1:02d}/{i % 12 + 1:02d}/{2016 + i % 10}'
        manager = f'Đơn vị mẫu {i % 100 + 1:03d}'
        branch = f'CN{i % 10 + 1:02d}'
        serial = f'DEMO-{i:08d}' if i % 113 else f'DEMO-REVIEW-{i % 3}'
        sheet.append([i + 1, name, f'{code:06d}', 'm' if kind == 'Vật tư' else 'cái', kind, 'DEMO', serial, 20 if kind == 'Vật tư' else 1, manager, '', '', f'{i % 100 + 1:04d}', branch, branch, f'DEMO-{i % 100}', manager, f'{i % 20:06d}', f'Nhân sự giả {i % 20:02d}', used, status])
    original = io.BytesIO()
    workbook.save(original)
    reference = Workbook(write_only=True)
    ref = reference.create_sheet('Chi_tiet')
    ref.append(['DỮ LIỆU GIẢ — KHÔNG PHẢI DANH MỤC NỘI BỘ'])
    ref.append(['Tên VTTB gốc','Loại VTTB','Nhóm lớn','Trạng thái rà soát','Ghi chú phân loại'])
    types = ['Máy in','Máy tính','Điều hòa','Vật tư văn phòng']
    for code in range(1,201):
        review = code % 19 == 0
        ref.append([f'Thiết bị mẫu {code:04d}',types[code % 4],'Nhóm DEMO','Cần rà soát' if review else 'Đã phân loại theo tên','Chỉ là dữ liệu giả' if review else ''])
    classification = io.BytesIO()
    reference.save(classification)
    return original.getvalue(), classification.getvalue()
