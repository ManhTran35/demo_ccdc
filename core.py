"""Data rules shared by the local app, HTML build, and verification tests."""
from __future__ import annotations

import gzip
import hashlib
import io
import json
import math
import re
import sqlite3
import uuid
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from difflib import get_close_matches
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
SAMPLE_CUTOFF = "2026-10-13"
UNKNOWN = "Chưa xác định"
UNCLASSIFIED = "Chưa phân loại"
AGE_GROUPS = ["Dưới 12 tháng", "12–35 tháng", "36–59 tháng", "60–71 tháng", "72–95 tháng", "Từ 96 tháng trở lên"]
DETAIL_COLUMNS = ["Chi nhánh", "Đơn vị quản lý", "Mã VTTB", "Tên VTTB", "Loại VTTB", "Tình trạng", "Số lượng", "ĐVT", "Ngày đưa vào sử dụng", "Số tháng sử dụng", "Nhóm thời gian sử dụng"]
REQUIRED = ["Tên VTTB", "Mã VTTB", "ĐVT", "Chủng loại", "Serial", "Số lượng", "Đơn vị quản lý", "Ngày đưa vào sử dụng", "Tình trạng"]


def vn_today() -> date:
    return datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).date()


def norm(value) -> str:
    return re.sub(r"\s+", " ", str(value) if value is not None else "").strip().casefold()


def clean(value) -> str:
    return re.sub(r"\s+", " ", str(value) if value is not None else "").strip()


def parse_date(value):
    if value is None or clean(value) == "":
        return None, "Thiếu ngày"
    if isinstance(value, datetime):
        return value.date(), "Hợp lệ"
    if isinstance(value, date):
        return value, "Hợp lệ"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if math.isfinite(value) and 1 <= value <= 2958465:
            try:
                return (datetime(1899, 12, 30) + timedelta(days=value)).date(), "Hợp lệ"
            except (OverflowError, ValueError):
                pass
        return None, "Ngày không hợp lệ"
    text = clean(value)
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date(), "Hợp lệ"
        except ValueError:
            pass
    return None, "Ngày không hợp lệ"


def full_months(start: date | None, today: date) -> int | None:
    if start is None or start > today:
        return None
    return (today.year - start.year) * 12 + today.month - start.month - (today.day < start.day)


def age_group(months: int | None) -> str:
    if months is None or pd.isna(months):
        return UNKNOWN
    for bound, label in zip((12, 36, 60, 72, 96), AGE_GROUPS):
        if months < bound:
            return label
    return AGE_GROUPS[-1]


def _cell_value(cell, header):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (int, float)) and any(x in header for x in ("Mã ", "Serial", "Bưu cục")):
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        fmt = cell.number_format or ""
        if re.fullmatch("0+", fmt):
            return str(value).zfill(len(fmt))
        return str(value)
    return value


def read_table(content: bytes, classification=False):
    """Read values without evaluating formulas or dropping duplicate source rows."""
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook["Chi_tiet"] if classification and "Chi_tiet" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows()
        header = None
        header_row = 0
        for line, cells in enumerate(rows, 1):
            values = [clean(c.value) for c in cells]
            if (("Tên VTTB gốc" in values and "Loại VTTB" in values) if classification else ("Tên VTTB" in values and "Mã VTTB" in values)):
                header = values
                header_row = line
                break
            if line >= 30:
                break
        if header is None:
            raise ValueError("Không tìm thấy dòng tiêu đề trong 30 dòng đầu của Excel.")
        seen = Counter()
        columns = []
        for idx, name in enumerate(header, 1):
            name = name or f"Cột {idx} (không tên)"
            seen[name] += 1
            columns.append(name if seen[name] == 1 else f"{name} [{seen[name]}]")
        result = []
        source_lines = []
        for line, cells in enumerate(rows, header_row + 1):
            row = [_cell_value(c, columns[i]) for i, c in enumerate(cells[:len(columns)])]
            row += [""] * (len(columns) - len(row))
            if all(v == "" for v in row):
                continue
            result.append(row)
            source_lines.append(line)
        return {"columns": columns, "rows": result, "source_lines": source_lines, "sheet": sheet.title}
    finally:
        workbook.close()


def classification_map(table):
    columns = table["columns"]
    required = ["Tên VTTB gốc", "Loại VTTB", "Nhóm lớn", "Trạng thái rà soát"]
    missing = [c for c in required if c not in columns]
    if missing:
        raise ValueError("Bảng phân loại thiếu cột: " + ", ".join(missing))
    index = {c: columns.index(c) for c in columns}
    buckets = {}
    for row in table["rows"]:
        key = norm(row[index["Tên VTTB gốc"]])
        if not key:
            continue
        record = [clean(row[index[c]]) for c in required[1:]]
        record.append(clean(row[index["Ghi chú phân loại"]]) if "Ghi chú phân loại" in index else "")
        record[0] = record[0] or UNCLASSIFIED
        buckets.setdefault(key, []).append(record)
    result = {}
    conflicts = 0
    for key, records in buckets.items():
        types = {(r[0], r[1]) for r in records}
        if len(types) > 1:
            result[key] = [UNCLASSIFIED, UNKNOWN, "Cần rà soát (xung đột)", "Một tên có nhiều loại/nhóm tham chiếu; chưa tự gán."]
            conflicts += 1
        else:
            review = any("rà soát" in norm(r[2]) or not r[2] or r[0] == UNCLASSIFIED for r in records)
            result[key] = [records[0][0], records[0][1] or UNKNOWN, "Cần rà soát" if review else "Đã phân loại", "; ".join(dict.fromkeys(r[3] for r in records if r[3]))]
    return result, {"reference_rows": len(table["rows"]), "unique_names": len(result), "conflicts": conflicts, "duplicate_reference_rows": sum(len(r) - 1 for r in buckets.values())}


def normalize(table, mapping, today=None):
    today = today or vn_today()
    columns = table["columns"]
    missing = [c for c in REQUIRED if c not in columns]
    if not any(c in columns for c in ("Chi nhánh mới", "Chi nhánh")):
        missing.append("Chi nhánh mới (hoặc Chi nhánh)")
    if missing:
        raise ValueError("File thiếu cột bắt buộc: " + ", ".join(missing))
    raw = pd.DataFrame(table["rows"], columns=columns)
    df = pd.DataFrame(index=raw.index)
    for col in columns:
        df[f"Gốc · {col}"] = raw[col]
    for col in ("Mã VTTB", "Tên VTTB", "ĐVT", "Chủng loại", "Serial", "Tình trạng", "Mã đơn vị", "Mã nhân viên"):
        df[col] = raw[col].map(clean) if col in raw else ""
    df["Chi nhánh"] = raw["Chi nhánh mới" if "Chi nhánh mới" in raw else "Chi nhánh"].map(clean).replace("", UNKNOWN)
    manager_col = "Đơn vị quản lý [2]" if "Đơn vị quản lý [2]" in raw else "Đơn vị quản lý"
    df["Đơn vị quản lý"] = raw[manager_col].map(clean)
    if manager_col != "Đơn vị quản lý":
        df["Đơn vị quản lý"] = df["Đơn vị quản lý"].where(df["Đơn vị quản lý"] != "", raw["Đơn vị quản lý"].map(clean))
    df["Đơn vị quản lý"] = df["Đơn vị quản lý"].replace("", UNKNOWN)
    # Canonical status strings only when the source differs by whitespace/case.
    canonical = {norm(s): s for s in ["Đang sử dụng", "Hỏng", "Tốt", "Tiêu hao đã xử lý", "Xuất cho đơn vị", "Thiết bị", "Vật tư"]}
    for col in ("Tình trạng", "Chủng loại"):
        df[col] = df[col].map(lambda x: canonical.get(norm(x), x or UNKNOWN))
    qty = pd.to_numeric(raw["Số lượng"].map(lambda x: clean(x) if isinstance(x, str) else x), errors="coerce")
    finite = qty.notna() & qty.map(lambda x: math.isfinite(x) if pd.notna(x) else False)
    invalid_qty = ~finite | (qty < 0)
    df["Số lượng"] = qty.where(finite)
    parsed_cache = {}
    def parsed(v):
        key = (type(v).__name__, str(v))
        if key not in parsed_cache:
            parsed_cache[key] = parse_date(v)
        return parsed_cache[key]
    parsed_dates = raw["Ngày đưa vào sử dụng"].map(parsed)
    dates = parsed_dates.map(lambda x: x[0])
    date_status = parsed_dates.map(lambda x: "Ngày ở tương lai" if x[0] and x[0] > today else x[1])
    df["Ngày đưa vào sử dụng"] = dates.map(lambda x: x.isoformat() if x else "")
    df["Trạng thái ngày"] = date_status
    df["Số tháng sử dụng"] = pd.array(dates.map(lambda x: full_months(x, today)), dtype="Int64")
    df["Nhóm thời gian sử dụng"] = df["Số tháng sử dụng"].map(age_group)
    mapped = df["Tên VTTB"].map(lambda x: mapping.get(norm(x), [UNCLASSIFIED, UNKNOWN, UNCLASSIFIED, ""]))
    for i, col in enumerate(["Loại VTTB", "Nhóm lớn", "Trạng thái phân loại", "Ghi chú phân loại"]):
        df[col] = mapped.map(lambda x: x[i])
    serial = df["Serial"].map(lambda x: norm(x).lstrip("'"))
    duplicate_serial = serial.ne("") & serial.duplicated(keep=False)
    raw_dedup = raw.drop(columns=["STT"], errors="ignore").astype(str)
    duplicate_rows = raw_dedup.duplicated(keep=False)
    flags = [[] for _ in range(len(df))]
    flag_masks = [(date_status != "Hợp lệ", date_status), (invalid_qty, "Số lượng không hợp lệ"), (duplicate_serial, "Serial nghi trùng"), (duplicate_rows, "Dòng nghi trùng"), (df["Loại VTTB"] == UNCLASSIFIED, UNCLASSIFIED), (df["Trạng thái phân loại"].str.contains("rà soát"), "Phân loại cần rà soát"), (df["Mã VTTB"] == "", "Thiếu mã VTTB"), (df["Tên VTTB"] == "", "Thiếu tên VTTB")]
    for mask, label in flag_masks:
        for idx in df.index[mask]:
            flags[idx].append(label.loc[idx] if isinstance(label, pd.Series) else label)
    df["Vấn đề dữ liệu"] = ["; ".join(f) for f in flags]
    df["Dòng Excel"] = table.get("source_lines", list(range(2, len(df) + 2)))
    df["_search"] = df[["Mã VTTB", "Tên VTTB", "Serial", "Đơn vị quản lý", "Chi nhánh"]].astype(str).agg(" ".join, axis=1).map(norm)
    review_mask = (df["Loại VTTB"] == UNCLASSIFIED) | df["Trạng thái phân loại"].str.contains("rà soát")
    report = {"source_rows": len(raw), "joined_rows": len(df), "invalid_quantity": int(invalid_qty.sum()), "missing_date": int((date_status == "Thiếu ngày").sum()), "invalid_date": int((date_status == "Ngày không hợp lệ").sum()), "future_date": int((date_status == "Ngày ở tương lai").sum()), "duplicate_serial_rows": int(duplicate_serial.sum()), "duplicate_serial_values": int(serial[duplicate_serial].nunique()), "duplicate_rows": int(duplicate_rows.sum()), "unclassified_rows": int((df["Loại VTTB"] == UNCLASSIFIED).sum()), "unclassified_names": int(df.loc[df["Loại VTTB"] == UNCLASSIFIED, "Tên VTTB"].map(norm).nunique()), "review_rows": int(review_mask.sum()), "review_names": int(df.loc[review_mask, "Tên VTTB"].map(norm).nunique()), "missing_code": int((df["Mã VTTB"] == "").sum()), "missing_name": int((df["Tên VTTB"] == "").sum()), "today": today.isoformat()}
    report["errors"] = (["File không có bản ghi dữ liệu."] if len(df) == 0 else []) + ([f"{report['invalid_quantity']} dòng có số lượng thiếu, âm hoặc không phải số."] if report["invalid_quantity"] else [])
    assert len(df) == len(raw), "Ghép phân loại đã làm thay đổi số dòng nguồn."
    return df, report


def refresh_dynamic_fields(df, today=None):
    """Recompute only daily age/date flags from a normalized Parquet snapshot.

    Classification and source normalization happen once during ingestion. Raw
    dates, identifiers, row order and static quality flags are never rewritten.
    """
    today = today or vn_today()
    out = df.copy()
    date_cache = {}
    for value in out["Ngày đưa vào sử dụng"].fillna("").unique():
        date_cache[value] = parse_date(value)[0]
    dates = out["Ngày đưa vào sử dụng"].fillna("").map(date_cache)
    previous = out["Trạng thái ngày"]
    statuses = [
        ("Ngày ở tương lai" if value > today else "Hợp lệ") if value
        else ("Thiếu ngày" if old == "Thiếu ngày" else "Ngày không hợp lệ")
        for value, old in zip(dates, previous)
    ]
    out["Trạng thái ngày"] = statuses
    out["Số tháng sử dụng"] = pd.array([full_months(v, today) for v in dates], dtype="Int64")
    out["Nhóm thời gian sử dụng"] = out["Số tháng sử dụng"].map(age_group)
    day_flags = {"Thiếu ngày", "Ngày không hợp lệ", "Ngày ở tương lai"}
    out["Vấn đề dữ liệu"] = [
        "; ".join(([status] if status != "Hợp lệ" else []) + [f for f in str(flags or "").split("; ") if f and f not in day_flags])
        for flags, status in zip(out["Vấn đề dữ liệu"].fillna(""), statuses)
    ]
    if "_search" not in out:
        out["_search"] = out[["Mã VTTB", "Tên VTTB", "Serial", "Đơn vị quản lý", "Chi nhánh"]].astype(str).agg(" ".join, axis=1).map(norm)
    return out


def apply_filters(df, branches=(), managers=(), types=(), statuses=(), search="", issue=""):
    mask = pd.Series(True, index=df.index)
    for col, selected in [("Chi nhánh", branches), ("Đơn vị quản lý", managers), ("Loại VTTB", types), ("Tình trạng", statuses)]:
        if selected:
            mask &= df[col].isin(selected)
    if clean(search):
        keys = ["Mã VTTB", "Tên VTTB", "Serial", "Đơn vị quản lý", "Chi nhánh"]
        haystack = df["_search"] if "_search" in df else df[keys].astype(str).agg(" ".join, axis=1).map(norm)
        mask &= haystack.str.contains(norm(search), regex=False)
    if issue:
        mask &= df["Vấn đề dữ liệu"].str.contains(issue, regex=False)
    return df.loc[mask]


def metrics(df):
    device = df["Chủng loại"].eq("Thiết bị")
    broken = device & df["Tình trạng"].eq("Hỏng")
    return {"rows": len(df), "codes": int(df.loc[df["Mã VTTB"].ne(""), "Mã VTTB"].nunique()), "using_codes": int(df.loc[df["Tình trạng"].eq("Đang sử dụng") & df["Mã VTTB"].ne(""), "Mã VTTB"].nunique()), "material_missing_codes": int(df.loc[df["Chủng loại"].eq("Vật tư") & df["Trạng thái ngày"].eq("Thiếu ngày") & df["Mã VTTB"].ne(""), "Mã VTTB"].nunique()), "device_rows": int(device.sum()), "broken_rows": int(broken.sum()), "broken_rate": float(broken.sum() / device.sum() * 100) if device.sum() else None}


def overview_tables(df):
    coded = df[df["Mã VTTB"].ne("")]
    by_branch = coded.groupby(["Chi nhánh", "Tình trạng"])["Mã VTTB"].nunique().unstack(fill_value=0)
    by_status = coded.groupby("Tình trạng")["Mã VTTB"].nunique().rename("Số mã duy nhất")
    using = df[df["Tình trạng"].eq("Đang sử dụng")]
    ages = using.loc[using["Số tháng sử dụng"].notna(), "Nhóm thời gian sử dụng"].value_counts().reindex(AGE_GROUPS, fill_value=0)
    return by_branch, by_status, ages, int(using["Số tháng sử dụng"].isna().sum())


def review_names(df, mapping):
    mask = df["Loại VTTB"].eq(UNCLASSIFIED) | df["Trạng thái phân loại"].str.contains("rà soát")
    names = df.loc[mask, ["Tên VTTB", "Loại VTTB", "Trạng thái phân loại", "Ghi chú phân loại"]].copy()
    names["_key"] = names["Tên VTTB"].map(norm)
    return names.groupby("_key", sort=False).agg(**{
        "Tên VTTB": ("Tên VTTB", "first"), "Loại VTTB": ("Loại VTTB", "first"),
        "Trạng thái phân loại": ("Trạng thái phân loại", "first"),
        "Ghi chú phân loại": ("Ghi chú phân loại", "first"), "Số dòng nguồn": ("Tên VTTB", "size")
    }).reset_index(drop=True)


def suggestions(name, mapping, limit=3):
    return get_close_matches(norm(name), list(mapping), n=limit, cutoff=.55)


def export_frame(df, columns, cutoff, today):
    out = df[list(columns)].copy()
    if "Số tháng sử dụng" in out:
        out["Số tháng sử dụng"] = out["Số tháng sử dụng"].astype(object).where(out["Số tháng sử dụng"].notna(), UNKNOWN)
    out["Ngày chốt dữ liệu"] = str(cutoff)
    out["Ngày tính tuổi sử dụng"] = today.isoformat() if isinstance(today, date) else str(today)
    return out


def _safe_csv(value):
    # Prevent spreadsheet formula evaluation while preserving normal string identifiers.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_csv(frame):
    return frame.map(_safe_csv).to_csv(index=False).encode("utf-8-sig")


def export_excel(frame):
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="xlsxwriter", engine_kwargs={"options": {"strings_to_numbers": False, "strings_to_formulas": False, "strings_to_urls": False}}) as writer:
        frame.to_excel(writer, sheet_name="Du_lieu_da_loc", index=False)
        sheet = writer.sheets["Du_lieu_da_loc"]
        sheet.freeze_panes(1, 0)
        sheet.autofilter(0, 0, len(frame), len(frame.columns) - 1)
        sheet.set_column(0, len(frame.columns) - 1, 24)
    return out.getvalue()


class VersionStore:
    """Each publication is a full immutable snapshot; restoration publishes a copy."""
    def __init__(self, root=None):
        self.root = Path(root or ROOT / "data" / "runtime")
        self.root.mkdir(parents=True, exist_ok=True)
        (self.root / "originals").mkdir(exist_ok=True)
        (self.root / "snapshots").mkdir(exist_ok=True)
        self.db = self.root / "versions.sqlite3"
        with sqlite3.connect(self.db) as conn:
            conn.execute("CREATE TABLE IF NOT EXISTS versions (id TEXT PRIMARY KEY, cutoff TEXT NOT NULL, name TEXT NOT NULL, created_at TEXT NOT NULL, published_at TEXT, digest TEXT NOT NULL, note TEXT NOT NULL)")

    def list(self, published_only=False):
        with sqlite3.connect(self.db) as conn:
            conn.row_factory = sqlite3.Row
            return [dict(r) for r in conn.execute("SELECT * FROM versions" + (" WHERE published_at IS NOT NULL" if published_only else "") + " ORDER BY COALESCE(published_at, created_at) DESC")]

    def create(self, content, name, cutoff, table, mapping, note=""):
        date.fromisoformat(str(cutoff))
        version_id = uuid.uuid4().hex
        digest = hashlib.sha256(content).hexdigest()
        original = self.root / "originals" / (digest + ".xlsx")
        if not original.exists():
            original.write_bytes(content)
        payload = {"table": table, "mapping": mapping}
        (self.root / "snapshots" / (version_id + ".json.gz")).write_bytes(gzip.compress(json.dumps(payload, ensure_ascii=False, separators=(",", ":"), allow_nan=False).encode()))
        with sqlite3.connect(self.db) as conn:
            conn.execute("INSERT INTO versions VALUES (?, ?, ?, ?, NULL, ?, ?)", (version_id, str(cutoff), name, datetime.now(timezone.utc).isoformat(), digest, note))
        return version_id

    def load(self, version_id):
        if not re.fullmatch(r"[a-f0-9]{32}", version_id):
            raise ValueError("Mã phiên không hợp lệ.")
        return json.loads(gzip.decompress((self.root / "snapshots" / (version_id + ".json.gz")).read_bytes()))

    def publish(self, version_id, role):
        if role != "Người duyệt":
            raise PermissionError("Chỉ vai trò Người duyệt được công bố (vai trò mô phỏng).")
        payload = self.load(version_id)
        _, report = normalize(payload["table"], payload["mapping"])
        if report["errors"]:
            raise ValueError("Không thể công bố: " + " ".join(report["errors"]))
        with sqlite3.connect(self.db) as conn:
            conn.execute("UPDATE versions SET published_at=? WHERE id=? AND published_at IS NULL", (datetime.now(timezone.utc).isoformat(), version_id))

    def latest(self):
        versions = self.list(True)
        return versions[0] if versions else None

    def as_of(self, cutoff):
        versions = [v for v in self.list(True) if v["cutoff"] <= str(cutoff)]
        return max(versions, key=lambda v: (v["cutoff"], v["published_at"])) if versions else None

    def restore(self, version_id, role):
        if role != "Người duyệt":
            raise PermissionError("Chỉ Người duyệt được khôi phục.")
        source = next(v for v in self.list(True) if v["id"] == version_id)
        payload = self.load(version_id)
        new_id = self.create((self.root / "originals" / (source["digest"] + ".xlsx")).read_bytes(), source["name"], source["cutoff"], payload["table"], payload["mapping"], f"Khôi phục từ phiên {version_id[:8]}")
        self.publish(new_id, role)
        return new_id


def seed_store(store):
    if store.list():
        return
    from services.demo import synthetic_files
    content, classification = synthetic_files(1000)
    table = read_table(content)
    mapping, _ = classification_map(read_table(classification, True))
    version_id = store.create(content, "inventory_demo.xlsx", SAMPLE_CUTOFF, table, mapping, "DỮ LIỆU GIẢ cho compatibility demo local.")
    store.publish(version_id, "Người duyệt")
