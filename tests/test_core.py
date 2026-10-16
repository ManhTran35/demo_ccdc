import csv
from datetime import date
import io
import tempfile
import unittest
from openpyxl import Workbook, load_workbook
import pandas as pd

from core import *


COLS = ['STT','Tên VTTB','Mã VTTB','ĐVT','Chủng loại','Serial','Số lượng','Đơn vị quản lý','Chi nhánh mới','Ngày đưa vào sử dụng','Tình trạng','Mã đơn vị']


def fixture():
    return {'columns':COLS.copy(),'rows':[
        [1,'  Máy   in ','0007','cái','Thiết bị','0001',1,'D1','A','20/03/2025','Đang sử dụng','001'],
        [2,'Vật tư lạ','0008','m','Vật tư','0002',2,'D2','A','','Tốt','002'],
        [3,'MÁY IN','0007','cái','Thiết bị','0001',1,'D3','B','20/03/2027','Hỏng','003'],
        [4,'Máy cần rà soát','0009','chiếc','Thiết bị','0003',1,'D4','B','31/02/2024','Đang sử dụng','004'],
    ],'source_lines':[2,3,4,5],'sheet':'Sheet1'}


MAPPING={'máy in':['Máy in','CNTT','Đã phân loại',''],'máy cần rà soát':['Máy tính','CNTT','Cần rà soát','Kiểm tra tên']}


class AgeTests(unittest.TestCase):
    def test_anniversary(self):
        self.assertEqual(full_months(date(2025,3,20),date(2026,3,19)),11)
        self.assertEqual(full_months(date(2025,3,20),date(2026,3,20)),12)

    def test_all_group_boundaries(self):
        for month, expected in [(0,0),(11,0),(12,1),(35,1),(36,2),(59,2),(60,3),(71,3),(72,4),(95,4),(96,5),(150,5)]:
            with self.subTest(month=month):
                self.assertEqual(age_group(month),AGE_GROUPS[expected])
                start_year=2026-month//12
                start_month=3-month%12
                if start_month<=0:start_month+=12;start_year-=1
                start=date(start_year,start_month,20)
                self.assertEqual(full_months(start,date(2026,3,20)),month)
                if month:self.assertEqual(full_months(start,date(2026,3,19)),month-1)

    def test_unknown_future(self):
        self.assertIsNone(full_months(None,date(2026,3,20)))
        self.assertIsNone(full_months(date(2027,3,20),date(2026,3,20)))
        self.assertEqual(age_group(None),UNKNOWN)

    def test_date_parsing(self):
        for value, expected in [('01/02/2024',date(2024,2,1)),('2024-02-29',date(2024,2,29)),(45292,date(2024,1,1)),(date(2025,1,2),date(2025,1,2))]:
            self.assertEqual(parse_date(value)[0],expected)
        self.assertEqual(parse_date('31/02/2024')[1],'Ngày không hợp lệ')
        self.assertEqual(parse_date('')[1],'Thiếu ngày')
        self.assertEqual(parse_date(0)[1],'Ngày không hợp lệ')

    def test_source_dates_are_always_day_first_including_time(self):
        for value, expected in [('01/02/2026',date(2026,2,1)),('02/01/2026',date(2026,1,2)),
                                ('13/10/2026',date(2026,10,13)),('29/02/2024 13:45:01.123456',date(2024,2,29)),
                                ('1/2/2026 0:00',date(2026,2,1)),('2026-02-01T00:00:00.123456',date(2026,2,1))]:
            with self.subTest(value=value): self.assertEqual(parse_date(value)[0],expected)
        for value in ['10/13/2026','29/02/2025','31/04/2026','01/02/26','01/02/2026 25:00']:
            with self.subTest(value=value): self.assertEqual(parse_date(value)[1],'Ngày không hợp lệ')

    def test_excel_text_native_and_serial_dates_preserve_calendar_day(self):
        wb=Workbook();ws=wb.active;ws.append(COLS)
        for value in ['01/02/2026', '02/01/2026', datetime(2026,2,1,12,30), 46054]:
            row=fixture()['rows'][0].copy();row[9]=value;ws.append(row)
        out=io.BytesIO();wb.save(out)
        frame,_=normalize(read_table(out.getvalue()),MAPPING,date(2026,10,15))
        self.assertEqual(frame['Ngày đưa vào sử dụng'].tolist(),['2026-02-01','2026-01-02','2026-02-01','2026-02-01'])
        self.assertEqual(display_frame(frame)['Ngày đưa vào sử dụng'].tolist(),['01/02/2026','02/01/2026','01/02/2026','01/02/2026'])
        self.assertEqual(frame.iloc[0]['Gốc · Ngày đưa vào sử dụng'],'01/02/2026')

    def test_day_refresh(self):
        before,_=normalize(fixture(),MAPPING,date(2026,3,19))
        after,_=normalize(fixture(),MAPPING,date(2026,3,20))
        self.assertEqual(before.iloc[0]['Số tháng sử dụng'],11)
        self.assertEqual(after.iloc[0]['Số tháng sử dụng'],12)


class DataTests(unittest.TestCase):
    def setUp(self):
        self.df,self.report=normalize(fixture(),MAPPING,date(2026,3,20))

    def test_exact_name_matching(self):
        self.assertEqual(self.df.iloc[0]['Loại VTTB'],'Máy in')
        self.assertEqual(self.df.iloc[2]['Loại VTTB'],'Máy in')
        self.assertEqual(self.df.iloc[1]['Loại VTTB'],UNCLASSIFIED)

    def test_keep_original_and_rows(self):
        self.assertEqual(len(self.df),4)
        self.assertEqual(self.df.iloc[0]['Gốc · Tên VTTB'],'  Máy   in ')
        self.assertEqual(self.report['source_rows'],self.report['joined_rows'])
        self.assertEqual(self.report['duplicate_serial_rows'],2)
        self.assertEqual(self.report['review_names'],2)

    def test_review_names_normalized_unique(self):
        table=fixture();table['rows'][3][1]=' VẬT   TƯ LẠ '
        df,report=normalize(table,MAPPING,date(2026,3,20))
        names=review_names(df,MAPPING)
        self.assertEqual(report['review_names'],1)
        self.assertEqual(len(names),1)
        self.assertEqual(names.iloc[0]['Số dòng nguồn'],2)

    def test_reference_duplicates_and_conflicts(self):
        table={'columns':['Tên VTTB gốc','Loại VTTB','Nhóm lớn','Trạng thái rà soát'],'rows':[['Máy in','Máy in','CNTT','Đã phân loại'],[' MÁY IN ','Máy in','CNTT','Đã phân loại']]}
        mapping,report=classification_map(table)
        self.assertEqual(report['duplicate_reference_rows'],1)
        df,_=normalize(fixture(),mapping,date(2026,3,20))
        self.assertEqual(len(df),4)
        table['rows'].append(['máy in','Máy tính','CNTT','Đã phân loại'])
        mapping,report=classification_map(table)
        self.assertEqual(report['conflicts'],1)
        self.assertEqual(mapping['máy in'][0],UNCLASSIFIED)
        self.assertIn('xung đột',mapping['máy in'][2])
        self.assertEqual(len(normalize(fixture(),mapping)[0]),4)

    def test_invalid_quantity_blocks(self):
        for value in ['',-1,'abc','0x11',float('inf')]:
            table=fixture();table['rows'][0][6]=value
            _,r=normalize(table,MAPPING,date(2026,3,20))
            self.assertTrue(r['errors'],repr(value))
        table=fixture();table['rows'][0][6]=0
        self.assertFalse(normalize(table,MAPPING)[1]['errors'])

    def test_missing_columns(self):
        table=fixture();table['columns'][2]='Wrong'
        with self.assertRaisesRegex(ValueError,'Mã VTTB'):normalize(table,MAPPING)

    def test_kpis_and_group_counts(self):
        m=metrics(self.df)
        self.assertEqual(m['using_codes'],2)
        self.assertEqual(m['material_missing_codes'],1)
        self.assertEqual(m['device_rows'],3)
        self.assertEqual(m['broken_rows'],1)
        self.assertAlmostEqual(m['broken_rate'],100/3)
        branch,status,ages,excluded=overview_tables(self.df)
        self.assertEqual(branch.loc['A','Đang sử dụng'],1)
        self.assertEqual(branch.loc['B','Hỏng'],1)
        self.assertEqual(status.loc['Đang sử dụng'],2)
        self.assertEqual(ages.sum(),1)
        self.assertEqual(excluded,1)

    def test_filters_share_same_rows(self):
        for args,expected in [({'branches':['A']},2),({'branches':['A'],'managers':['D1']},1),({'types':[UNCLASSIFIED]},1),({'statuses':['Hỏng']},1),({'search':'0007'},2),({'search':'máy in'},2),({'issue':'Serial nghi trùng'},2),({'branches':['A'],'managers':['D3']},0)]:
            filtered=apply_filters(self.df,**args)
            self.assertEqual(len(filtered),expected)
            self.assertEqual(metrics(filtered)['rows'],expected)
            self.assertEqual(len(export_frame(filtered,DETAIL_COLUMNS,'2026-03-20',date(2026,3,20))),expected)

    def test_status_excludes_age(self):
        _,_,ages,excluded=overview_tables(apply_filters(self.df,statuses=['Hỏng']))
        self.assertEqual(ages.sum(),0)
        self.assertEqual(excluded,0)

    def test_csv_and_excel_preserve_strings_metadata(self):
        filtered=apply_filters(self.df,branches=['A'])
        frame=export_frame(filtered,DETAIL_COLUMNS+['Mã đơn vị','Serial'],'2026-03-20',date(2026,3,20))
        csv_rows=list(csv.DictReader(io.StringIO(export_csv(frame).decode('utf-8-sig'))))
        self.assertEqual(len(csv_rows),2)
        self.assertEqual(csv_rows[0]['Mã VTTB'],'0007')
        self.assertEqual(csv_rows[0]['Mã đơn vị'],'001')
        self.assertEqual(csv_rows[0]['Serial'],'0001')
        self.assertEqual(csv_rows[1]['Số tháng sử dụng'],UNKNOWN)
        self.assertEqual(csv_rows[0]['Ngày tính tuổi sử dụng'],'20/03/2026')
        self.assertEqual(csv_rows[0]['Ngày đưa vào sử dụng'],'20/03/2025')
        wb=load_workbook(io.BytesIO(export_excel(frame)),read_only=True)
        values=list(wb.active.values)
        self.assertEqual(len(values),3)
        self.assertEqual(values[1][2],'0007')
        self.assertEqual(values[1][-2: ],('20/03/2026','20/03/2026'))
        wb.close()

    def test_excel_numeric_identifier_format(self):
        wb=Workbook();ws=wb.active;ws.append(COLS);ws.append(fixture()['rows'][0]);ws.cell(2,3).value=7;ws.cell(2,3).number_format='0000';b=io.BytesIO();wb.save(b)
        self.assertEqual(read_table(b.getvalue())['rows'][0][2],'0007')

    def test_duplicate_manager_headers(self):
        table=fixture();table['columns'].append('Đơn vị quản lý [2]')
        for row in table['rows']:row.append('D_NEW')
        df,_=normalize(table,MAPPING)
        self.assertTrue((df['Đơn vị quản lý']=='D_NEW').all())
        self.assertEqual(df.iloc[0]['Gốc · Đơn vị quản lý'],'D1')


class VersionTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.store=VersionStore(self.temp.name)
        self.first=self.store.create(b'original bytes','test.xlsx','2026-03-20',fixture(),MAPPING)
    def tearDown(self):self.temp.cleanup()

    def test_role_and_pending(self):
        self.assertIsNone(self.store.latest())
        with self.assertRaises(PermissionError):self.store.publish(self.first,'Người nhập')
        self.assertIsNone(self.store.latest())
        self.store.publish(self.first,'Người duyệt')
        self.assertEqual(self.store.latest()['id'],self.first)

    def test_bad_file_does_not_replace(self):
        self.store.publish(self.first,'Người duyệt')
        bad=fixture();bad['rows'][0][6]=-5
        second=self.store.create(b'bad original','bad.xlsx','2026-04-01',bad,MAPPING)
        with self.assertRaises(ValueError):self.store.publish(second,'Người duyệt')
        self.assertEqual(self.store.latest()['id'],self.first)

    def test_date_selection_no_future(self):
        self.store.publish(self.first,'Người duyệt')
        self.assertIsNone(self.store.as_of('2026-03-19'))
        self.assertEqual(self.store.as_of('2026-03-20')['id'],self.first)
        second=self.store.create(b'next original','next.xlsx','2026-04-01',fixture(),MAPPING)
        self.store.publish(second,'Người duyệt')
        self.assertEqual(self.store.as_of('2026-03-31')['id'],self.first)
        self.assertEqual(self.store.as_of('2026-04-01')['id'],second)
        self.assertEqual(len(self.store.load(second)['table']['rows']),4)

    def test_restore_immutable_snapshot(self):
        self.store.publish(self.first,'Người duyệt')
        before=self.store.load(self.first)
        new_id=self.store.restore(self.first,'Người duyệt')
        self.assertNotEqual(new_id,self.first)
        self.assertEqual(self.store.load(new_id),before)
        self.assertEqual(len(self.store.list(True)),2)
        self.assertEqual(self.store.latest()['id'],new_id)
        self.assertEqual(self.store.as_of('2026-03-20')['id'],new_id)
        self.assertEqual(len(list((self.store.root/'originals').glob('*'))),1)


if __name__=='__main__':unittest.main()
