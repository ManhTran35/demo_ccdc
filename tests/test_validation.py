"""Synthetic-only security and canonical-transform regression checks."""
from dataclasses import replace
from datetime import date
import io
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

from core import classification_map, metrics, normalize, read_table
from services.config import Settings
from services.transform import parse_classification, refresh_dynamic_fields, transform_inventory
from services.validation import ValidationError, validate_workbook

COLUMNS = ["Tên VTTB", "Mã VTTB", "ĐVT", "Chủng loại", "Serial", "Số lượng", "Đơn vị quản lý", "Chi nhánh mới", "Ngày đưa vào sử dụng", "Tình trạng", "Mã đơn vị"]
ROWS = [
    ["  Máy   mẫu ", "0007", "cái", "Thiết bị", "0001", 1, "Đơn vị giả A", "Chi nhánh giả", "15/10/2025", "Đang sử dụng", "001"],
    ["Vật tư mẫu", "0008", "m", "Vật tư", "0002", 2, "Đơn vị giả B", "Chi nhánh giả", "", "Tốt", "002"],
    ["MÁY MẪU", "0007", "cái", "Thiết bị", "0001", 1, "Đơn vị giả A", "Chi nhánh giả", "16/10/2026", "Hỏng", "001"],
]
MAP_COLUMNS = ["Tên VTTB gốc", "Loại VTTB", "Nhóm lớn", "Trạng thái rà soát"]
MAP_ROWS = [["Máy mẫu", "Máy giả", "Nhóm giả", "Đã phân loại"], ["Vật tư mẫu", "Vật tư giả", "Nhóm giả", "Đã phân loại"]]


def workbook_bytes(columns=None, rows=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(COLUMNS if columns is None else columns)
    for row in ROWS if rows is None else rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def rewrite_zip(content, updates=None, extras=None):
    output = io.BytesIO()
    updates = updates or {}
    with ZipFile(io.BytesIO(content)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename in updates:
                action = updates[item.filename]
                data = action(data) if callable(action) else action
            target.writestr(item.filename, data)
        for key, value in (extras or {}).items():
            target.writestr(key, value)
    return output.getvalue()


class WorkbookSecurityTests(unittest.TestCase):
    def setUp(self):
        self.content = workbook_bytes()
        self.settings = Settings()

    def test_normal_workbook_allowed(self):
        self.assertIsNone(validate_workbook(self.content, "source.xlsx", self.settings))

    def test_fake_xlsx_and_unsupported_types_rejected(self):
        for data, name in [(b"not a zip", "x.xlsx"), (self.content, "x.xlsm"), (self.content, "x.xls"), (b"", "x.xlsx")]:
            with self.subTest(name=name, length=len(data)), self.assertRaises(ValidationError):
                validate_workbook(data, name, self.settings)

    def test_formula_cells_rejected_including_hidden_worksheet(self):
        workbook = Workbook()
        workbook.active.append(COLUMNS)
        workbook.active.append(ROWS[0])
        hidden = workbook.create_sheet("hidden")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "=1+1"
        output = io.BytesIO()
        workbook.save(output)
        with self.assertRaises(ValidationError) as raised:
            validate_workbook(output.getvalue(), "x.xlsx")
        self.assertEqual(raised.exception.code, "formula_not_allowed")
        self.assertIn("values-only", str(raised.exception))
        self.assertEqual(raised.exception.report['filename'], 'x.xlsx')
        self.assertEqual(raised.exception.report['sheet'], 'hidden')
        self.assertEqual(raised.exception.report['cell'], 'A1')

    def test_filter_print_metadata_and_formatting_dropdown_formulas_are_not_cell_values(self):
        workbook = load_workbook(io.BytesIO(self.content))
        sheet = workbook.active
        sheet.auto_filter.ref = 'A1:K4'
        sheet.print_area = 'A1:K4'
        sheet.print_title_rows = '1:1'
        sheet.conditional_formatting.add('A2:K4', FormulaRule(formula=['A2="example"']))
        validation = DataValidation(type='list', formula1='"A,B"'); validation.add('A2:A4')
        sheet.add_data_validation(validation)
        output = io.BytesIO(); workbook.save(output); workbook.close()
        validate_workbook(output.getvalue(), 'metadata.xlsx')
        self.assertEqual(len(transform_inventory(output.getvalue(), {}, date(2026,10,15)).dataframe), 3)

    def test_arbitrary_or_dynamic_builtin_defined_names_remain_rejected(self):
        for name, expression in [('custom', 'Sheet!$A$1'), ('_xlnm._FilterDatabase', 'OFFSET(Sheet!$A$1,0,0)'),
                                  ('_xlnm.Print_Area', "'[other.xlsx]Sheet'!$A$1")]:
            workbook = load_workbook(io.BytesIO(self.content))
            workbook.defined_names.add(DefinedName(name, attr_text=expression))
            output = io.BytesIO(); workbook.save(output); workbook.close()
            with self.subTest(name=name), self.assertRaises(ValidationError) as caught:
                validate_workbook(output.getvalue(), 'names.xlsx')
            self.assertEqual(caught.exception.report['defined_name'], name)
            self.assertNotIn(expression, str(caught.exception))

    def test_cached_formula_result_never_reaches_table_parser(self):
        content = rewrite_zip(self.content, {'xl/worksheets/sheet1.xml': lambda b: b.replace(
            b'<c r="B2" t="inlineStr"><is><t>0007</t></is></c>', b'<c r="B2"><f>1+1</f><v>999</v></c>')})
        from unittest.mock import patch
        with patch('core.read_table', side_effect=AssertionError('cached value reached Pandas')) as parser:
            with self.assertRaises(ValidationError) as caught:
                transform_inventory(content, {}, date(2026,10,15), filename='cached.xlsx')
        parser.assert_not_called()
        self.assertEqual(caught.exception.report['cell'], 'B2')

    def test_macro_external_link_dde_rejected(self):
        for name, content in [("xl/vbaProject.bin", b"macro"),
                              ("xl/externalLinks/externalLink1.xml", b"<x/>"),
                              ("xl/embeddings/oleObject1.bin", b"ole"),
                              ("xl/connections.xml", b"<connections/>"),
                              ("xl/dde.xml", b"<root><ddeLink/></root>")]:
            with self.subTest(name=name), self.assertRaises(ValidationError):
                validate_workbook(rewrite_zip(self.content, extras={name: content}), "x.xlsx")

    def test_external_relationship_rejected(self):
        relationship = b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="hyperlink" Target="https://example.invalid/private" TargetMode="External"/></Relationships>'
        content = rewrite_zip(self.content, extras={"xl/worksheets/_rels/sheet1.xml.rels": relationship})
        with self.assertRaises(ValidationError) as raised:
            validate_workbook(content, "x.xlsx")
        self.assertEqual(raised.exception.code, "external_relationship")

    def test_dtd_entities_rejected_utf8_and_utf16(self):
        text = '<?xml version="1.0"?><!DOCTYPE r [<!ENTITY private "SENSITIVE">]><r>&private;</r>'
        for encoding in ("utf-8", "utf-16"):
            with self.subTest(encoding=encoding), self.assertRaises(ValidationError) as raised:
                validate_workbook(rewrite_zip(self.content, extras={"docProps/untrusted.xml": text.encode(encoding)}), "x.xlsx")
            self.assertNotIn("SENSITIVE", str(raised.exception))

    def test_path_traversal_and_duplicate_members_rejected(self):
        for filename in ("../escape.xml", "/escape.xml", "xl\\escape.xml", "C:/escape.xml", "xl/./escape.xml", "xl//escape.xml", "xl/workbook.xml"):
            with self.subTest(filename=filename), self.assertRaises(ValidationError) as raised:
                validate_workbook(rewrite_zip(self.content, extras={filename: b"<x/>"}), "x.xlsx")
            self.assertEqual(raised.exception.code, "unsafe_zip_path")

    def test_upload_uncompressed_and_member_limits(self):
        for settings in (replace(self.settings, max_upload_mb=0), replace(self.settings, max_uncompressed_mb=0)):
            with self.subTest(settings=settings), self.assertRaises(ValidationError):
                validate_workbook(self.content, "x.xlsx", settings)
        many = {f"extra/{i}.txt": b"x" for i in range(4096)}
        with self.assertRaises(ValidationError) as raised:
            validate_workbook(rewrite_zip(self.content, extras=many), "x.xlsx")
        self.assertEqual(raised.exception.code, "member_limit")

    def test_dimension_and_actual_cell_limit(self):
        for settings in (replace(self.settings, max_columns=2), replace(self.settings, max_rows=1)):
            # Actual max_rows is checked after canonical header discovery.
            if settings.max_rows == 1:
                with self.assertRaises(ValidationError):
                    transform_inventory(self.content, {}, settings=settings)
            else:
                with self.assertRaises(ValidationError):
                    validate_workbook(self.content, "x.xlsx", settings)
        content = rewrite_zip(self.content, {"xl/worksheets/sheet1.xml": lambda b: b.replace(b'r="A2"', b'r="ZZ2"')})
        with self.assertRaises(ValidationError):
            validate_workbook(content, "x.xlsx")

    def test_malformed_xml_rejected(self):
        with self.assertRaises(ValidationError):
            validate_workbook(rewrite_zip(self.content, {"xl/worksheets/sheet1.xml": b"<broken>"}), "x.xlsx")


class TransformTests(unittest.TestCase):
    def setUp(self):
        self.content = workbook_bytes()
        self.mapping, _ = parse_classification(workbook_bytes(MAP_COLUMNS, MAP_ROWS))
        self.today = date(2026, 10, 15)

    def test_parquet_preserves_rows_raw_strings_ids_and_kpi(self):
        transformed = transform_inventory(self.content, self.mapping, self.today, Settings())
        actual = pd.read_parquet(io.BytesIO(transformed.parquet_bytes))
        expected, _ = normalize(read_table(self.content), self.mapping, self.today)
        self.assertEqual(len(actual), len(ROWS))
        self.assertEqual(metrics(actual), metrics(expected))
        self.assertEqual(actual["Mã VTTB"].tolist(), ["0007", "0008", "0007"])
        self.assertEqual(actual["Serial"].tolist(), ["0001", "0002", "0001"])
        self.assertEqual(actual.iloc[0]["Gốc · Tên VTTB"], ROWS[0][0])
        self.assertTrue(all(isinstance(value, str) for column in actual if column.startswith("Gốc · ") for value in actual[column]))
        self.assertEqual(actual["Dòng Excel"].tolist(), [2, 3, 4])
        self.assertTrue(transformed.report["valid"])
        self.assertEqual(transformed.report["row_count"], len(ROWS))
        self.assertEqual(transformed.metrics["units"], 2)
        self.assertEqual(len(transformed.metrics["quantity_by_unit"]), 2)

    def test_mixed_raw_columns_safe_for_arrow(self):
        rows = [r.copy() for r in ROWS]
        rows[0][5], rows[1][5] = 1, "2"
        result = transform_inventory(workbook_bytes(rows=rows), self.mapping, self.today)
        self.assertEqual(pd.read_parquet(io.BytesIO(result.parquet_bytes))["Gốc · Số lượng"].tolist(), ["1", "2", "1"])

    def test_numeric_identifier_number_format_survives(self):
        workbook = Workbook(); sheet = workbook.active
        sheet.append(COLUMNS); sheet.append(ROWS[0])
        sheet.cell(2, 2, 7).number_format = "0000"
        output = io.BytesIO(); workbook.save(output)
        result = transform_inventory(output.getvalue(), self.mapping, self.today)
        self.assertEqual(result.dataframe.iloc[0]["Mã VTTB"], "0007")

    def test_negative_missing_and_non_numeric_quantity_block(self):
        for quantity in (-1, "", "abc"):
            rows = [r.copy() for r in ROWS]; rows[0][5] = quantity
            with self.subTest(quantity=quantity), self.assertRaises(ValidationError) as raised:
                transform_inventory(workbook_bytes(rows=rows), self.mapping, self.today)
            self.assertFalse(raised.exception.report["valid"])
            self.assertEqual(raised.exception.report["invalid_quantity"], 1)

    def test_missing_schema_and_empty_file_block(self):
        for content in (workbook_bytes(["Tên VTTB", "Mã VTTB"], [["Fake", "0001"]]), workbook_bytes(rows=[])):
            with self.assertRaises(ValidationError):
                transform_inventory(content, self.mapping, self.today)

    def test_colliding_duplicate_header_labels_are_reported(self):
        columns = COLUMNS + ["Đơn vị quản lý [2]", "Đơn vị quản lý"]
        rows = [row + ["Fake", "Fake"] for row in ROWS]
        with self.assertRaises(ValidationError) as raised:
            transform_inventory(workbook_bytes(columns, rows), self.mapping, self.today)
        self.assertEqual(raised.exception.code, "duplicate_columns")

    def test_warning_thresholds_and_unknown_domains_preserve_rows(self):
        rows = [r.copy() for r in ROWS]; rows[0][9] = "Trạng thái giả mới"
        result = transform_inventory(workbook_bytes(rows=rows), self.mapping, self.today,
                                     replace(Settings(), warning_limits={"missing_date": 1.0}))
        self.assertEqual(len(result.dataframe), len(ROWS))
        codes = [warning["code"] for warning in result.report["warnings"]]
        self.assertIn("unknown_status", codes)
        self.assertNotIn("missing_date", codes)

    def test_reference_conflict_does_not_multiply_or_deduplicate_rows(self):
        mapping, report = parse_classification(workbook_bytes(MAP_COLUMNS, MAP_ROWS + [["MÁY MẪU", "Khác", "Khác", "Đã phân loại"]]))
        result = transform_inventory(self.content, mapping, self.today)
        self.assertEqual(report["conflicts"], 1)
        self.assertEqual(len(result.dataframe), 3)
        self.assertEqual(result.dataframe.iloc[0]["Loại VTTB"], "Chưa phân loại")

    def test_daily_refresh_matches_canonical_rules_without_reclassification(self):
        result = transform_inventory(self.content, self.mapping, self.today)
        actual = refresh_dynamic_fields(result.dataframe, date(2026, 10, 16))
        expected, _ = normalize(read_table(self.content), self.mapping, date(2026, 10, 16))
        for column in ("Số tháng sử dụng", "Nhóm thời gian sử dụng", "Trạng thái ngày", "Vấn đề dữ liệu"):
            self.assertEqual(actual[column].astype(object).fillna("NULL").astype(str).tolist(), expected[column].astype(object).fillna("NULL").astype(str).tolist())
        self.assertEqual(actual["Loại VTTB"].tolist(), result.dataframe["Loại VTTB"].tolist())
        self.assertEqual(actual["Gốc · Tên VTTB"].tolist(), result.dataframe["Gốc · Tên VTTB"].tolist())
        self.assertEqual(result.dataframe.iloc[2]["Trạng thái ngày"], "Ngày ở tương lai")


if __name__ == "__main__":
    unittest.main()
