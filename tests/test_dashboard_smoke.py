"""UI and repository checks against synthetic data, without network/cloud calls."""
from datetime import date
import io
from openpyxl import load_workbook
from pathlib import Path
import os
import runpy
import tempfile
import unittest
from unittest.mock import patch

import streamlit as st
from streamlit.testing.v1 import AppTest

from core import metrics, vn_today, refresh_dynamic_fields
from services.auth import Identity
from services.config import Settings
from services.demo import synthetic_files
from services.repository import DataRepository
from services.storage import LocalStorageAdapter
from services.transform import parse_classification, transform_inventory
from services.drive import DriveClient

ROOT = Path(__file__).resolve().parents[1]


class DashboardSmokeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {"APP_ENV":"local", "STORAGE_BACKEND":"local", "LOCAL_DATA_DIR":self.temp.name}, clear=False)
        self.env.start()
        self.kservice = os.environ.pop("K_SERVICE", None)
        st.cache_resource.clear()
        self.repo = DataRepository(Settings.from_env())
        self.repo.seed_synthetic_demo()

    def tearDown(self):
        st.cache_resource.clear()
        self.env.stop()
        if self.kservice is not None:
            os.environ["K_SERVICE"] = self.kservice
        self.temp.cleanup()

    def app(self):
        app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60).run()
        self.assertFalse(app.exception, [e.message for e in app.exception])
        return app

    def cloud_env(self):
        return patch.dict(os.environ, {"APP_ENV":"staging", "STORAGE_BACKEND":"gcs", "GCP_PROJECT_ID":"test-project", "GCS_BUCKET":"test-bucket", "IAP_AUDIENCE":"/test-verified-audience", "ADMIN_EMAILS":"admin@example.com"})

    def test_local_kpis_details_admin_and_filters(self):
        app = self.app()
        frame = self.repo.read_active_dataset(vn_today())
        expected = metrics(frame)
        self.assertEqual(app.metric[0].value, f"{expected['using_codes']:,}".replace(",", "."))
        app.sidebar.multiselect[0].set_value(["CN01"]).run()
        self.assertFalse(app.exception)
        self.assertTrue(app.sidebar.multiselect[1].options)
        manager = app.sidebar.multiselect[1].options[0]
        app.sidebar.multiselect[1].set_value([manager]).run()
        app.sidebar.multiselect[0].set_value(["CN02"]).run()
        self.assertEqual(app.sidebar.multiselect[1].value, [])
        app.sidebar.date_input[0].set_value(date(2026,10,12)).run()
        self.assertTrue(any("Chưa có phiên" in item.value for item in app.info))
        app.sidebar.date_input[0].set_value(date(2026,10,13)).run()
        app.sidebar.radio[0].set_value("Chi tiết dữ liệu").run()
        self.assertFalse(app.exception)
        self.assertTrue(app.dataframe)
        self.assertFalse(any(c.startswith('_') for c in app.multiselect[-1].options))
        app.sidebar.radio[0].set_value("Quản trị").run()
        self.assertFalse(app.exception)
        self.assertTrue(any("mô phỏng" in i.value for i in app.info))

    def test_community_entry_point_is_synthetic_read_only_and_preserves_filters_exports(self):
        with tempfile.TemporaryDirectory() as directory, patch('services.config.tempfile.gettempdir', return_value=directory):
            app = AppTest.from_file(str(ROOT / 'streamlit_app.py'), default_timeout=60).run()
            self.assertFalse(app.exception, [e.message for e in app.exception])
            self.assertEqual(app.sidebar.radio[0].options, ['Tổng quan', 'Chi tiết dữ liệu'])
            self.assertFalse(any(s.label == 'Vai trò mô phỏng' for s in app.selectbox))
            self.assertEqual(len(app.metric), 3)
            self.assertTrue(any('1,000 / 1,000' in x.value for x in app.markdown))
            self.assertTrue(any('dữ liệu giả' in x.value for x in app.info))
            app.sidebar.multiselect[0].set_value(['CN01']).run()
            self.assertTrue(any('100 / 1,000' in x.value for x in app.markdown))
            app.sidebar.radio[0].set_value('Chi tiết dữ liệu').run()
            self.assertFalse(app.exception)
            next(b for b in app.button if 'Chuẩn bị CSV' in b.label).click().run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.get('download_button')), 2)
            self.assertFalse(app.get('file_uploader'))
            self.assertEqual(len(self.repo.read_active_dataset(vn_today())), 1000)


    def test_cloud_has_no_admin_upload_or_simulated_role_and_no_excel_parse(self):
        with self.cloud_env(), patch('services.repository.create_storage', return_value=LocalStorageAdapter(self.temp.name)), \
             patch('services.auth.verify_iap_identity', return_value=Identity('viewer@example.com','subject',False)), \
             patch('core.read_table', side_effect=AssertionError('must not parse Excel on read')), \
             patch('core.normalize', side_effect=AssertionError('must not normalize source on read')):
            app = self.app()
            self.assertEqual(app.sidebar.radio[0].options, ['Tổng quan','Chi tiết dữ liệu'])
            self.assertFalse(any(s.label == 'Vai trò mô phỏng' for s in app.selectbox))
            self.assertEqual(len(app.metric), 3)

    def test_drive_entrypoint_refresh_and_historical_date_use_named_snapshots(self):
        from test_drive import FakeDrive, remote, FILE_B
        client = FakeDrive(rows=30)
        with tempfile.TemporaryDirectory() as directory, patch('services.config.tempfile.gettempdir', return_value=directory), \
                patch('streamlit.secrets', {'data_source': 'google_drive'}), \
                patch.object(DriveClient, 'from_service_account', return_value=client):
            app = AppTest.from_file(str(ROOT / 'streamlit_app.py'), default_timeout=60).run()
            self.assertFalse(app.exception, [e.message for e in app.exception])
            self.assertFalse(app.error)
            self.assertEqual(len(app.metric), 3)
            self.assertEqual(app.sidebar.radio[0].options, ['Tổng quan', 'Chi tiết dữ liệu'])
            self.assertTrue(any('Ton_CCDC_13102026' in c.value for c in app.caption))
            content = synthetic_files(8)[0]
            client.contents[FILE_B] = content
            client.files.append(remote(FILE_B, 'Ton_CCDC_14102026.xlsx', content=content))
            next(b for b in app.button if b.label == 'Làm mới dữ liệu').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.sidebar.date_input[0].value, date(2026,10,14))
            self.assertTrue(any('8 / 8 bản ghi' in m.value for m in app.markdown))
            app.sidebar.date_input[0].set_value(date(2026,10,13)).run()
            self.assertFalse(app.exception)
            self.assertTrue(any('30 / 30 bản ghi' in m.value for m in app.markdown))
            app.sidebar.radio[0].set_value('Chi tiết dữ liệu').run()
            table = app.dataframe[0].value
            self.assertEqual(table.iloc[1]['Ngày đưa vào sử dụng'], '02/02/2017')

    def test_configured_drive_without_access_shows_error_instead_of_demo(self):
        from services.drive import DriveError
        with patch('streamlit.secrets', {'data_source':'google_drive'}), \
                patch.object(DriveClient, 'from_service_account', side_effect=DriveError('Chưa có quyền đọc Drive.')):
            app=AppTest.from_file(str(ROOT/'streamlit_app.py'), default_timeout=60).run()
            self.assertFalse(app.exception)
            self.assertFalse(app.metric)
            self.assertTrue(any('quyền đọc Drive' in e.value for e in app.error))

    def test_cloud_without_iap_fails_before_opening_repository(self):
        with self.cloud_env(), patch('services.repository.create_storage') as create:
            app = self.app()
            self.assertTrue(any('IAP' in item.value for item in app.error))
            self.assertFalse(app.metric)
            create.assert_not_called()

    def test_cloud_storage_error_is_sanitized(self):
        storage = LocalStorageAdapter(self.temp.name)
        with self.cloud_env(), patch('services.repository.create_storage', return_value=storage), \
             patch('services.auth.verify_iap_identity', return_value=Identity('viewer@example.com','subject',False)), \
             patch.object(storage, 'get', side_effect=RuntimeError('PRIVATE_BUCKET_TOKEN_DO_NOT_DISPLAY')):
            app = self.app()
            self.assertTrue(app.error)
            self.assertNotIn('PRIVATE_BUCKET_TOKEN', ' '.join(x.value for x in app.error))
            self.assertFalse(app.metric)

    def test_active_change_and_rollback_refresh_default_dataset(self):
        app = self.app()
        app.sidebar.multiselect[0].set_value(["CN01"]).run()
        old = self.repo.get_active_version()
        inventory, reference = synthetic_files(12)
        workbook = load_workbook(io.BytesIO(inventory))
        for row in workbook.active.iter_rows(min_row=2):
            row[13].value = "NEW_BRANCH"
        buffer = io.BytesIO()
        workbook.save(buffer)
        inventory = buffer.getvalue()
        mapping, _ = parse_classification(reference)
        result = transform_inventory(inventory,mapping)
        manager = self.repo.manager
        stage = manager.stage_upload(inventory,'synthetic.xlsx','local-admin')
        preview = manager.prepare_version(stage,result.parquet_bytes,mapping,result.report,result.metrics,'2026-10-14','local-admin',old['_generation'])
        pointer = manager.commit(preview['id'],preview['confirmation_token'],'local-admin',old['_generation'])
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(app.sidebar.date_input[0].value,date(2026,10,14))
        self.assertEqual(app.sidebar.multiselect[0].value, [])
        self.assertEqual(len(self.repo.read_active_dataset(vn_today())),12)
        rollback = manager.preview_rollback(old['active_version'],pointer['_generation'])
        manager.rollback(old['active_version'],'local-admin',pointer['_generation'],rollback['confirmation_token'])
        app.run()
        self.assertEqual(app.sidebar.date_input[0].value,date(2026,10,13))
        self.assertEqual(self.repo.select_version('2099-01-01')['id'],old['active_version'])

    def test_daily_age_changes_without_reclassifying_or_mutating_cache(self):
        frame = self.repo.read_active_dataset(date(2026,3,19))
        old = frame.copy(deep=True)
        with patch('core.normalize', side_effect=AssertionError('no re-normalization')):
            refreshed = refresh_dynamic_fields(frame,date(2026,3,20))
        self.assertTrue(frame.equals(old))
        self.assertTrue((refreshed['Loại VTTB']==old['Loại VTTB']).all())
        self.assertTrue((refreshed['Gốc · Ngày đưa vào sử dụng']==old['Gốc · Ngày đưa vào sử dụng']).all())

    def test_scheduled_refresh_error_forces_safe_full_app_boundary(self):
        class RerunSignal(BaseException):
            pass
        from unittest.mock import Mock
        failing = Mock()
        failing.get_active_version.side_effect = RuntimeError('PRIVATE_TOKEN_MUST_NOT_APPEAR')
        with patch('streamlit.fragment', side_effect=lambda **kwargs: lambda fn: fn):
            module = runpy.run_path(str(ROOT / 'app.py'), run_name='ccdc_fragment_test')
        with patch('streamlit.rerun', side_effect=RerunSignal) as rerun, \
             self.assertLogs('ccdc_fragment_test', level='ERROR') as log:
            with self.assertRaises(RerunSignal):
                module['refresh_state']('2026-10-15','version',1,failing)
        rerun.assert_called_once_with(scope='app')
        self.assertNotIn('PRIVATE_TOKEN', ' '.join(log.output))


if __name__ == '__main__':
    unittest.main()
